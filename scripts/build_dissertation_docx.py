"""Build the final dissertation DOCX from docs/dissertation.md."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


TITLE = "Design and Implementation of an Explainable Machine Learning Model for Financial Fraud Detection"
INK = "17202A"
MID = "44546A"
LIGHT = "D9E1E8"
PALE = "F4F6F8"


def set_font(run, name="Georgia", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, value=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), LIGHT)
        borders.append(element)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(table.rows):
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, MID if row_index == 0 else (PALE if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, size=8.5, bold=(row_index == 0), color="FFFFFF" if row_index == 0 else INK)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_inline(paragraph, text):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(paragraph.add_run(part[2:-2]), bold=True)
        elif part.startswith("`") and part.endswith("`"):
            set_font(paragraph.add_run(part[1:-1]), name="Consolas", size=9)
        elif part.startswith("*") and part.endswith("*"):
            set_font(paragraph.add_run(part[1:-1]), italic=True)
        else:
            set_font(paragraph.add_run(part))


def configure_document(document):
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (
        ("Title", 22, 0, 18),
        ("Heading 1", 16, 0, 10),
        ("Heading 2", 12.5, 12, 6),
        ("Heading 3", 11, 10, 4),
    ):
        style = document.styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("Explainable Machine Learning for Fraud Detection"), size=8, color=MID)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1")
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()

    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_cover(document):
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    add_inline(p, TITLE)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    set_font(p.add_run("MSc Dissertation"), size=15, bold=True, color=MID)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(185)
    set_font(p.add_run("September 2026"), size=11)
    document.add_page_break()


def parse_table(lines, start, document):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([p.strip() for p in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", p) for p in rows[1]):
        rows.pop(1)
    table = document.add_table(rows=1, cols=len(rows[0]))
    for j, value in enumerate(rows[0]):
        table.rows[0].cells[j].text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
    style_table(table)
    return i


def add_figure(document, project_root, alt, relative_path):
    image_path = (project_root / "docs" / relative_path).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(5.85))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", alt)
        doc_pr[0].set("title", alt.split(".", 1)[0])
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    cap_run = caption.add_run(alt)
    set_font(cap_run, size=9, italic=True, color=MID)


def render_markdown(document, markdown_path, project_root):
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    i = 1
    abstract_seen = False
    toc_added = False
    references_mode = False
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if title == "Abstract":
                abstract_seen = True
            elif abstract_seen and not toc_added:
                document.add_page_break()
                toc_heading = document.add_paragraph()
                set_font(toc_heading.add_run("Table of Contents"), size=16, bold=True, color=INK)
                toc_heading.paragraph_format.space_after = Pt(10)
                toc = document.add_paragraph()
                add_field(toc, ' TOC \\o "1-2" \\h \\z \\u ', "Update field to display contents")
                document.add_page_break()
                toc_added = True
            references_mode = title == "References"
            heading = document.add_heading(title, level=1)
            if title.startswith("Chapter ") or title == "References" or title.startswith("Appendix"):
                heading.paragraph_format.page_break_before = True
            i += 1
            continue
        if stripped.startswith("### "):
            heading_text = stripped[4:].strip()
            if heading_text in {"Aim", "Objectives", "Application verification"}:
                spacer = document.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(1)
                set_font(spacer.add_run(" "), size=1)
            heading = document.add_heading(heading_text, level=2)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            heading.paragraph_format.left_indent = Cm(0)
            heading.paragraph_format.first_line_indent = Cm(0)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)
            i += 1
            continue
        image = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image:
            add_figure(document, project_root, image.group(1), image.group(2))
            i += 1
            continue
        if stripped.startswith("|"):
            i = parse_table(lines, i, document)
            continue
        if re.match(r"^\d+\. ", stripped):
            number, content = stripped.split(". ", 1)
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.65)
            set_font(p.add_run(f"{number}.\t"))
            add_inline(p, content)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue
        if stripped.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            i += 1
            continue
        if stripped.startswith("*Table ") and stripped.endswith("*"):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(stripped[1:-1]), size=9, italic=True, color=MID)
            i += 1
            continue
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if references_mode:
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.65)
            p.paragraph_format.space_after = Pt(7)
        else:
            previous = document.paragraphs[-2] if len(document.paragraphs) > 1 else None
            if previous is not None and previous.style.name == "Normal" and previous.text:
                p.paragraph_format.first_line_indent = Cm(0.6)
        add_inline(p, stripped)
        i += 1


def build(markdown_path, output, project_root):
    document = Document()
    configure_document(document)
    add_cover(document)
    render_markdown(document, markdown_path, project_root)
    document.core_properties.title = TITLE
    document.core_properties.subject = "MSc dissertation on explainable credit-card fraud detection"
    document.core_properties.keywords = "fraud detection, machine learning, random forest, SHAP, explainable AI"
    document.core_properties.comments = "Generated from the verified project evidence source."
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, default=Path("docs/dissertation.md"))
    parser.add_argument("--output", type=Path, default=Path("dissertation.docx"))
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    build(args.markdown.resolve(), args.output.resolve(), args.project_root.resolve())
    print(f"Saved {args.output.resolve()}")


if __name__ == "__main__":
    main()
